# -*- coding: utf-8 -*-
"""استيراد الطلاب من Excel مع تقرير أخطاء عربي مفصّل (مرحلة أساس التطبيق).

قواعد العمل:
- كل الصفوف تُدقَّق أولاً، ثم تُنشأ الحسابات في معاملة واحدة.
- أي خطأ في صف يمنع استيراده ويُدرج في التقرير برقم الصف والعمود والسبب.
- الشُعب تُنشأ تلقائياً عند غيابها (باسمها وقسمها ومرحلتها ونوع دراستها).
- كلمة المرور الأولية: الرقم الجامعي افتراضياً (أو عشوائية حسب الإعداد)
  مع إجبار التغيير عند أول دخول.
"""
import io
import re
import secrets
from dataclasses import dataclass, field
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.errors import ApiError
from app.models import AcademicYear, Department, Group, Student, Term, User
from app.models.enums import StudentStatus, StudyType, UserRole, UserStatus
from app.security import hash_password

AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")

HEADER_ALIASES: dict[str, set[str]] = {
    "university_id": {
        "university_id", "university id", "id",
        "الرقم الجامعي", "الرقم_الجامعي", "رقم الطالب", "رقم الجامعة",
    },
    "full_name_ar": {
        "full_name_ar", "full name", "name",
        "الاسم", "اسم الطالب", "الاسم الرباعي", "الاسم الثلاثي",
    },
    "department": {
        "department_code", "department", "department code",
        "القسم", "رمز القسم", "القسم_الرمز",
    },
    "stage": {"stage", "المرحلة", "مرحلة"},
    "study_type": {"study_type", "نوع الدراسة", "الدراسة", "نوع الدراسه"},
    "group": {"group_name", "group", "الشعبة", "شعبة"},
    "enrollment_year": {"enrollment_year", "سنة التسجيل", "سنة القبول"},
    "email": {"email", "البريد الإلكتروني", "البريد الالكتروني", "ايميل"},
    "phone": {"phone", "الهاتف", "رقم الهاتف", "الموبايل"},
}

REQUIRED_FIELDS = ("university_id", "full_name_ar", "department", "stage", "study_type")

MAX_ERRORS = 300
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
PHONE_RE = re.compile(r"^[0-9+\- ]{7,20}$")

STUDY_MAP = {
    "صباحي": StudyType.MORNING,
    "صباحا": StudyType.MORNING,
    "صباح": StudyType.MORNING,
    "morning": StudyType.MORNING,
    "am": StudyType.MORNING,
    "مسائي": StudyType.EVENING,
    "مساء": StudyType.EVENING,
    "مساءً": StudyType.EVENING,
    "evening": StudyType.EVENING,
    "pm": StudyType.EVENING,
}


@dataclass
class Issue:
    row: int
    column: str
    code: str
    message_ar: str

    def as_dict(self) -> dict:
        return {
            "row": self.row,
            "column": self.column,
            "code": self.code,
            "message_ar": self.message_ar,
        }


@dataclass
class ImportResult:
    rows_total: int = 0
    rows_valid: int = 0
    rows_invalid: int = 0
    created: list[dict] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)
    groups_created: list[str] = field(default_factory=list)
    truncated: bool = False


def _norm(value: object) -> str:
    """تنظيف قيمة الخلية: نص، أرقام عربية → لاتينية، وضغط المسافات."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).translate(AR_DIGITS).strip()
    return re.sub(r"\s+", " ", text)


def _header_map(header_row: tuple) -> tuple[dict[str, int], list[str]]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(header_row):
        key = _norm(raw).lower()
        if not key:
            continue
        for fieldname, aliases in HEADER_ALIASES.items():
            if key in {a.lower() for a in aliases} and fieldname not in mapping:
                mapping[fieldname] = index
                break
    missing = [name for name in REQUIRED_FIELDS if name not in mapping]
    return mapping, missing


def _parse_int(text: str) -> int | None:
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _resolve_department(
    raw: str, departments: list[Department]
) -> Department | None:
    key = raw.strip().lower()
    for dep in departments:
        if dep.code.lower() == key or dep.name_ar.strip() == raw.strip():
            return dep
    return None


def _ensure_term(db: Session) -> Term:
    """الفصل الحالي، وإن لم يوجد يُنشأ فصل افتراضي للسنة الحالية (تهيئة أولى)."""
    year = db.scalar(select(AcademicYear).where(AcademicYear.is_current.is_(True)))
    if year is None:
        year = db.scalar(select(AcademicYear).order_by(AcademicYear.label.desc()))
    if year is None:
        year = AcademicYear(label=str(datetime.now(timezone.utc).year), is_current=True)
        db.add(year)
        db.flush()

    term = db.scalar(
        select(Term).where(Term.academic_year_id == year.id, Term.is_current.is_(True))
    )
    if term is None:
        term = db.scalar(
            select(Term).where(Term.academic_year_id == year.id).order_by(Term.name)
        )
    if term is None:
        term = Term(
            academic_year_id=year.id,
            name="الفصل الأول",
            is_current=True,
        )
        db.add(term)
        db.flush()
    return term


def import_students(
    db: Session,
    content: bytes,
    *,
    dry_run: bool,
    actor: User,
) -> ImportResult:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - أي ملف تالف يعطي نفس الرسالة للمستخدم
        raise ApiError(
            400, "invalid_file", "تعذّر قراءة ملف Excel؛ تأكد أنه بصيغة xlsx"
        ) from exc

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        raise ApiError(400, "empty_file", "الملف فارغ")

    mapping, missing = _header_map(rows[0])
    if missing:
        raise ApiError(
            400,
            "invalid_template",
            "ملف Excel لا يحتوي كل الأعمدة المطلوبة",
            {"missing_columns": missing},
        )

    departments = list(db.scalars(select(Department).where(Department.deleted_at.is_(None))))
    term = _ensure_term(db)

    # نطاق رئيس القسم: قسمه فقط
    actor_department_id = None
    if actor.role == UserRole.DEPARTMENT_HEAD:
        from app.models import Staff  # استيراد محلي لتجنّب الدورة

        staff = db.get(Staff, actor.id)
        actor_department_id = staff.department_id if staff else None

    result = ImportResult()
    parsed: list[dict] = []
    seen_ids: set[str] = set()
    groups_cache: dict[tuple, Group] = {
        (g.department_id, g.stage, g.name.strip(), g.study_type, g.term_id): g
        for g in db.scalars(select(Group).where(Group.deleted_at.is_(None)))
    }

    for offset, raw_row in enumerate(rows[1:], start=2):
        values = {name: _norm(raw_row[idx]) for name, idx in mapping.items()}
        if not any(values.values()):
            continue  # صف فارغ تماماً

        result.rows_total += 1
        row_errors: list[Issue] = []

        university_id = values.get("university_id", "")
        full_name = values.get("full_name_ar", "")
        department_raw = values.get("department", "")
        stage_raw = values.get("stage", "")
        study_raw = values.get("study_type", "")

        if not university_id:
            row_errors.append(
                Issue(offset, "university_id", "missing_university_id", "الرقم الجامعي مطلوب")
            )
        elif not re.fullmatch(r"[0-9A-Za-z/_-]{3,20}", university_id):
            row_errors.append(
                Issue(offset, "university_id", "invalid_university_id", "الرقم الجامعي غير صالح")
            )
        elif university_id in seen_ids:
            row_errors.append(
                Issue(offset, "university_id", "duplicate_in_file", "الرقم الجامعي مكرر داخل الملف")
            )

        if len(full_name) < 3:
            row_errors.append(
                Issue(offset, "full_name_ar", "invalid_name", "اسم الطالب مطلوب (3 أحرف على الأقل)")
            )

        department = _resolve_department(department_raw, departments) if department_raw else None
        if department is None:
            row_errors.append(
                Issue(
                    offset,
                    "department",
                    "unknown_department",
                    "القسم غير معروف؛ استخدم رمز القسم (CS, IS, MS, CY) أو اسمه العربي",
                )
            )
        elif actor_department_id is not None and department.id != actor_department_id:
            row_errors.append(
                Issue(
                    offset,
                    "department",
                    "not_your_department",
                    "هذا الصف خارج قسمك؛ رئيس القسم يستورد طلاب قسمه فقط",
                )
            )

        stage = _parse_int(stage_raw)
        if stage is None or not 1 <= stage <= 6:
            row_errors.append(
                Issue(offset, "stage", "invalid_stage", "المرحلة يجب أن تكون رقماً من 1 إلى 6")
            )

        study_type = STUDY_MAP.get(study_raw.lower()) if study_raw else None
        if study_type is None:
            row_errors.append(
                Issue(
                    offset,
                    "study_type",
                    "invalid_study_type",
                    "نوع الدراسة يجب أن يكون صباحي أو مسائي",
                )
            )

        email = values.get("email", "")
        if email and not EMAIL_RE.fullmatch(email):
            row_errors.append(
                Issue(offset, "email", "invalid_email", "البريد الإلكتروني غير صالح")
            )

        phone = values.get("phone", "")
        if phone and not PHONE_RE.fullmatch(phone):
            row_errors.append(
                Issue(offset, "phone", "invalid_phone", "رقم الهاتف غير صالح")
            )

        enrollment_year = _parse_int(values.get("enrollment_year", ""))
        if values.get("enrollment_year") and (
            enrollment_year is None or not 1990 <= enrollment_year <= 2100
        ):
            row_errors.append(
                Issue(
                    offset,
                    "enrollment_year",
                    "invalid_enrollment_year",
                    "سنة التسجيل يجب أن تكون سنة بين 1990 و2100",
                )
            )

        if university_id and not row_errors:
            seen_ids.add(university_id)

        if row_errors:
            result.rows_invalid += 1
            result.errors.extend(row_errors)
            continue

        result.rows_valid += 1
        parsed.append(
            {
                "row": offset,
                "university_id": university_id,
                "full_name_ar": full_name,
                "department": department,
                "stage": stage,
                "study_type": study_type,
                "group_name": values.get("group", ""),
                "enrollment_year": enrollment_year,
                "email": email or None,
                "phone": phone or None,
        }
        )

    # التحقق من التكرار في قاعدة البيانات (بعد قراءة الملف كاملاً)
    ids = [item["university_id"] for item in parsed]
    existing: set[str] = set()
    if ids:
        existing = {
            row[0]
            for row in db.execute(
                select(User.university_id).where(User.university_id.in_(ids))
            )
        }
    if existing:
        still_valid: list[dict] = []
        for item in parsed:
            if item["university_id"] in existing:
                result.rows_valid -= 1
                result.rows_invalid += 1
                result.errors.append(
                    Issue(
                        item["row"],
                        "university_id",
                        "duplicate_university_id",
                        "يوجد حساب بهذا الرقم الجامعي مسبقاً",
                    )
                )
            else:
                still_valid.append(item)
        parsed = still_valid

    # ترتيب الأخطاء حسب رقم الصف
    result.errors.sort(key=lambda issue: (issue.row, issue.column))
    if len(result.errors) > MAX_ERRORS:
        result.errors = result.errors[:MAX_ERRORS]
        result.truncated = True

    if dry_run:
        db.rollback()
        return result

    # الإنشاء الفعلي — معاملة واحدة
    initial_mode = settings.initial_password_mode
    year_hint = datetime.now(timezone.utc).year
    try:
        for item in parsed:
            group_id = None
            group_name = item["group_name"]
            if group_name:
                key = (
                    item["department"].id,
                    item["stage"],
                    group_name,
                    item["study_type"],
                    term.id,
                )
                group = groups_cache.get(key)
                if group is None:
                    group = Group(
                        department_id=item["department"].id,
                        stage=item["stage"],
                        name=group_name,
                        study_type=item["study_type"],
                        term_id=term.id,
                    )
                    db.add(group)
                    db.flush()
                    groups_cache[key] = group
                    result.groups_created.append(
                        f"{item['department'].name_ar} / المرحلة {item['stage']} / {group_name}"
                    )
                group_id = group.id

            initial_password = (
                secrets.token_urlsafe(6)[:10]
                if initial_mode == "random"
                else item["university_id"]
            )

            user = User(
                university_id=item["university_id"],
                full_name_ar=item["full_name_ar"],
                role=UserRole.STUDENT,
                email=item["email"],
                phone=item["phone"],
                password_hash=hash_password(initial_password, rounds=settings.import_bcrypt_rounds),
                status=UserStatus.ACTIVE,
                must_change_password=True,
            )
            db.add(user)
            db.flush()

            db.add(
                Student(
                    user_id=user.id,
                    department_id=item["department"].id,
                    stage=item["stage"],
                    study_type=item["study_type"],
                    group_id=group_id,
                    enrollment_year=item["enrollment_year"] or year_hint,
                    status=StudentStatus.ACTIVE,
                )
            )
            created_row = {
                "university_id": item["university_id"],
                "full_name_ar": item["full_name_ar"],
            }
            if initial_mode == "random":
                created_row["initial_password"] = initial_password
            result.created.append(created_row)

        db.commit()
    except Exception:
        db.rollback()
        raise

    return result


def build_template() -> bytes:
    """قالب Excel جاهز بأعمدة معتمدة وصف مثال — لتوزيعه على الأقسام."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "students"
    headers = [
        "university_id",
        "full_name_ar",
        "department_code",
        "stage",
        "study_type",
        "group_name",
        "enrollment_year",
        "email",
        "phone",
    ]
    sheet.append(headers)
    sheet.append(
        ["20240001", "اسم الطالب الرباعي", "CS", 1, "صباحي", "A", 2026, "", "07700000000"]
    )
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
