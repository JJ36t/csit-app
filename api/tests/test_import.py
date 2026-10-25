# -*- coding: utf-8 -*-
"""اختبارات استيراد الطلاب من Excel: القالب، الفحص، الأخطاء، الإنشاء، ودخول الطالب."""
import io

from openpyxl import Workbook, load_workbook

HEADERS = [
    "university_id",
    "full_name_ar",
    "department_code",
    "stage",
    "study_type",
    "group_name",
    "enrollment_year",
    "email",
    "phone",
]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

VALID_ROWS = [
    ["20250001", "علي حسن محمد", "CS", 1, "صباحي", "A", 2026, "", ""],
    ["20250002", "زينب كريم جواد", "IS", 2, "مسائي", "B", 2025, "z@example.com", "07701234567"],
    ["20250003", "محمد صادق علي", "CY", 3, "صباحي", "", 2026, "", ""],
]


def _xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_template_download(admin_client):
    res = admin_client.get("/api/v1/admin/students/template")
    assert res.status_code == 200
    workbook = load_workbook(io.BytesIO(res.content))
    sheet = workbook.active
    assert sheet.cell(row=1, column=1).value == "university_id"


def test_dry_run_reports_errors_without_creating(admin_client):
    rows = VALID_ROWS + [
        ["", "بلا رقم جامعي", "CS", 1, "صباحي", "", 2026, "", ""],
        ["20250009", "بيانات خاطئة", "XX", 9, "غير معروف", "", 2026, "", ""],
        ["20250001", "مكرر داخل الملف", "CS", 1, "صباحي", "", 2026, "", ""],
    ]
    res = admin_client.post(
        "/api/v1/admin/students/import",
        params={"dry_run": "true"},
        files={"file": ("students.xlsx", _xlsx(rows), XLSX_MIME)},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["rows_total"] == 6
    assert body["rows_valid"] == 3
    assert body["rows_invalid"] == 3
    assert body["created_count"] == 0

    codes = {error["code"] for error in body["errors"]}
    assert {
        "missing_university_id",
        "unknown_department",
        "invalid_stage",
        "invalid_study_type",
        "duplicate_in_file",
    } <= codes

    # تأكيد أن الفحص لم يُنشئ أي حساب
    listing = admin_client.get("/api/v1/admin/students").json()
    assert listing["total"] == 0


def test_import_creates_accounts_and_student_can_login(client, admin_client):
    res = admin_client.post(
        "/api/v1/admin/students/import",
        params={"dry_run": "false"},
        files={"file": ("students.xlsx", _xlsx(VALID_ROWS), XLSX_MIME)},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created_count"] == 3
    assert len(body["groups_created"]) == 2  # الشعبة A والشعبة B

    listing = admin_client.get(
        "/api/v1/admin/students", params={"department_code": "CS"}
    ).json()
    assert listing["total"] == 1
    item = listing["items"][0]
    assert item["study_type"] == "morning"
    assert item["group_name"] == "A"
    assert item["department_name_ar"] == "علوم الحاسوب"

    # الطالب يدخل بكلمة المرور الأولية (رقمه الجامعي) ويُطلب منه تغييرها
    login = client.post(
        "/api/v1/auth/login",
        json={
            "university_id": "20250001",
            "password": "20250001",
            "device": {"fingerprint": "student-device-01", "model": "Redmi"},
        },
    )
    assert login.status_code == 200, login.text
    payload = login.json()
    assert payload["user"]["must_change_password"] is True
    assert payload["user"]["student"]["department"]["code"] == "CS"
    assert payload["user"]["student"]["group_name"] == "A"


def test_duplicate_account_reported_on_reimport(admin_client):
    files = {"file": ("s.xlsx", _xlsx(VALID_ROWS[:1]), XLSX_MIME)}
    first = admin_client.post(
        "/api/v1/admin/students/import", params={"dry_run": "false"}, files=files
    ).json()
    assert first["created_count"] == 1

    again = admin_client.post(
        "/api/v1/admin/students/import",
        params={"dry_run": "true"},
        files={"file": ("s.xlsx", _xlsx(VALID_ROWS[:1]), XLSX_MIME)},
    ).json()
    codes = {error["code"] for error in again["errors"]}
    assert "duplicate_university_id" in codes


def test_invalid_file_rejected(admin_client):
    res = admin_client.post(
        "/api/v1/admin/students/import",
        params={"dry_run": "true"},
        files={"file": ("students.txt", b"hello", "text/plain")},
    )
    assert res.status_code == 400
    assert res.json()["code"] == "invalid_file_type"


def test_missing_columns_rejected(admin_client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["university_id", "full_name_ar"])
    sheet.append(["20250010", "ناقص الأعمدة"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    res = admin_client.post(
        "/api/v1/admin/students/import",
        params={"dry_run": "true"},
        files={"file": ("s.xlsx", buffer.getvalue(), XLSX_MIME)},
    )
    assert res.status_code == 400
    body = res.json()
    assert body["code"] == "invalid_template"
    assert "stage" in body["details"]["missing_columns"]
