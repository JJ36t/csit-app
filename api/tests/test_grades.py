# -*- coding: utf-8 -*-
"""اختبارات الدرجات (N3): تقييمات، إدخال دفعي مدقَّق، نشر بموافقة، قراءة، وتصدير Excel."""
import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditLog, Notification
from world import auth_headers, build_world, login


def _create_assessment(
    client,
    headers,
    offering_id: str,
    *,
    name: str = "الامتحان النصفي",
    kind: str = "midterm",
    max_score: float = 50,
    weight: float = 25,
):
    return client.post(
        "/api/v1/assessments",
        json={
            "offering_id": offering_id,
            "name": name,
            "kind": kind,
            "max_score": max_score,
            "weight": weight,
        },
        headers=headers,
    )


def _enter(client, headers, assessment_id: str, rows):
    return client.post(
        f"/api/v1/assessments/{assessment_id}/grades",
        json={"rows": rows},
        headers=headers,
    )


def test_teacher_creates_assessment_and_student_is_forbidden(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-01"))
    student = auth_headers(login(client, "S-100", "grade-dev-02"))

    created = _create_assessment(client, teacher, world["offering"])
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["max_score"] == 50.0
    assert body["weight"] == 25.0
    assert body["is_published"] is False

    forbidden = _create_assessment(client, student, world["offering"])
    assert forbidden.status_code == 403
    assert forbidden.json()["code"] == "not_offering_staff"


def test_assessment_weights_cannot_exceed_100(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-03"))

    assert _create_assessment(client, teacher, world["offering"], weight=60).status_code == 200
    over = _create_assessment(client, teacher, world["offering"], weight=50)
    assert over.status_code == 400
    assert over.json()["code"] == "weight_exceeds_100"


def test_bulk_entry_is_audited_and_validated(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-04"))
    assessment_id = _create_assessment(client, teacher, world["offering"]).json()["assessment_id"]

    first = _enter(
        client,
        teacher,
        assessment_id,
        [
            {"student_id": world["student"], "score": 40},
            {"student_id": world["other"], "score": 45.5},
        ],
    )
    assert first.status_code == 200, first.text
    assert first.json()["created"] == 2

    second = _enter(
        client, teacher, assessment_id, [{"student_id": world["student"], "score": 42}]
    )
    assert second.json()["updated"] == 1

    with SessionLocal() as db:
        actions = {
            row.action
            for row in db.scalars(select(AuditLog).where(AuditLog.entity == "grade"))
        }
    assert "grade_entered" in actions
    assert "grade_updated" in actions

    over_max = _enter(
        client, teacher, assessment_id, [{"student_id": world["student"], "score": 51}]
    )
    assert over_max.status_code == 400
    assert over_max.json()["code"] == "score_exceeds_max"

    outsider = _enter(
        client, teacher, assessment_id, [{"student_id": world["outsider"], "score": 10}]
    )
    assert outsider.status_code == 400
    assert outsider.json()["code"] == "student_not_enrolled"


def test_published_assessment_rejects_grade_edits(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-05"))
    admin = auth_headers(login(client, "ADMIN-9", "grade-admin-01"))
    assessment_id = _create_assessment(client, teacher, world["offering"]).json()["assessment_id"]
    _enter(client, teacher, assessment_id, [{"student_id": world["student"], "score": 40}])

    assert (
        client.post(
            f"/api/v1/assessments/{assessment_id}/publish", headers=admin
        ).status_code
        == 200
    )

    blocked = _enter(
        client, teacher, assessment_id, [{"student_id": world["student"], "score": 45}]
    )
    assert blocked.status_code == 409
    assert blocked.json()["code"] == "assessment_published"


def test_student_sees_only_published_grades(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-06"))
    admin = auth_headers(login(client, "ADMIN-9", "grade-admin-02"))
    student = auth_headers(login(client, "S-100", "grade-dev-07"))

    assessment_id = _create_assessment(client, teacher, world["offering"]).json()["assessment_id"]
    _enter(client, teacher, assessment_id, [{"student_id": world["student"], "score": 40}])

    before = client.get("/api/v1/grades", headers=student)
    assert before.status_code == 200
    assert before.json()["items"][0]["assessments"] == []

    client.post(f"/api/v1/assessments/{assessment_id}/publish", headers=admin)

    after = client.get("/api/v1/grades", headers=student).json()["items"][0]
    assert len(after["assessments"]) == 1
    assert after["assessments"][0]["score"] == 40.0
    # 40/50 × وزن 25 = 20
    assert after["weighted_total"] == 20.0
    assert after["weight_sum"] == 25.0


def test_publish_requires_college_admin_and_notifies_students(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-08"))
    admin = auth_headers(login(client, "ADMIN-9", "grade-admin-03"))

    assessment_id = _create_assessment(client, teacher, world["offering"]).json()["assessment_id"]
    _enter(client, teacher, assessment_id, [{"student_id": world["student"], "score": 40}])

    forbidden = client.post(f"/api/v1/assessments/{assessment_id}/publish", headers=teacher)
    assert forbidden.status_code == 403
    assert forbidden.json()["code"] == "not_college_admin"

    published = client.post(f"/api/v1/assessments/{assessment_id}/publish", headers=admin)
    assert published.status_code == 200, published.text
    assert published.json()["notified"] == 2  # S-100 وS-200 المسجَّلان

    with SessionLocal() as db:
        notifications = list(
            db.scalars(select(Notification).where(Notification.type == "grades_published"))
        )
        audit = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "assessment_published",
                AuditLog.entity_id == assessment_id,
            )
        )
    assert len(notifications) == 2
    assert audit is not None


def test_offering_matrix_and_excel_export(client):
    world = build_world()
    teacher = auth_headers(login(client, "T-100", "grade-dev-09"))
    student = auth_headers(login(client, "S-100", "grade-dev-10"))
    assessment_id = _create_assessment(client, teacher, world["offering"]).json()["assessment_id"]
    _enter(
        client,
        teacher,
        assessment_id,
        [
            {"student_id": world["student"], "score": 40},
            {"student_id": world["other"], "score": 30},
        ],
    )

    matrix = client.get(f"/api/v1/offerings/{world['offering']}/grades", headers=teacher)
    assert matrix.status_code == 200
    body = matrix.json()
    assert len(body["assessments"]) == 1
    rows = {row["university_id"]: row for row in body["students"]}
    assert rows["S-100"]["scores"][assessment_id] == 40.0
    assert rows["S-100"]["weighted_total"] == 20.0
    assert rows["S-200"]["weighted_total"] == 15.0

    forbidden = client.get(f"/api/v1/offerings/{world['offering']}/grades", headers=student)
    assert forbidden.status_code == 403

    export = client.get(f"/api/v1/offerings/{world['offering']}/grades/export", headers=teacher)
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(io.BytesIO(export.content)).active
    assert sheet.max_row == 3  # عنوان + طالبان
    assert sheet.cell(row=1, column=1).value == "الرقم الجامعي"
